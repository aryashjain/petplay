"""Builds Pawdio.pptx from the captures made by capture.js.

Every element becomes its own picture with a PowerPoint entrance animation
(auto-play, timed like the HTML deck), GIFs stay animated inside circular
frames, and each slide gets a transition + speaker notes.

Usage: python build_pptx.py
"""
import json
import struct
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

HERE = Path(__file__).parent
BUILD = HERE / "build"
OUT = HERE.parent / "Pawdio.pptx"
PX = 6350  # EMU per CSS px on a 1920x1080 canvas (13.333in wide)

NOTES = [
    "Hook: music changes how dogs feel. Shelter studies found classical music meant more resting and less barking "
    "(Kogan et al., 2012), and soft rock / reggae lowered stress (Univ. of Glasgow & Scottish SPCA, 2017). "
    "Dogs also hear up to ~45 kHz, way beyond our ~20 kHz. So what if their favourite toy made the music?",
    "The idea: dogs love balls and they respond to music, so we fused them. Step 1, the dog just plays. "
    "Step 2, a motion sensor sealed in the ball feels every hit and spin 200 times a second. "
    "Step 3, a web page turns that motion into live piano music that always stays in key.",
    "Inside the ball: ESP32-C6 (RISC-V, Wi-Fi 6, BLE 5, deep sleep), a 6-axis MPU IMU on I2C "
    "(accelerometer + gyroscope), and a small sealed LiPo. Firmware samples at 200 Hz or more, applies a light "
    "low-pass filter, detects impacts from acceleration spikes, packs compact binary frames and streams them over "
    "a WebSocket. It auto-reconnects and sleeps when the ball is still, waking on the IMU motion interrupt.",
    "The music brain is plain HTML, CSS and JavaScript, no phone app. It decodes the binary frames, maps impact "
    "strength to note velocity and rotation speed to pitch and note density, snaps every note to the chosen scale "
    "so it never sounds harsh, and synthesises piano with the Web Audio API. It also has a live motion visualiser, "
    "a connection indicator, and scale, tempo and volume controls.",
    "Future scope: the same tiny module fits any toy. Plushies that sing when squeezed, tug ropes that crescendo, "
    "cat toys, health and activity insights for owners and vets, multiple toys jamming as a band, and sensory / "
    "therapy toys for kids and senior dogs.",
    "Thank you! Pawdio: every fetch is a song. Happy to take questions or show a live demo.",
]

# CSS entrance name -> (PowerPoint presetID, presetSubtype)
PRESETS = {
    "up": (42, 0), "left": (2, 8), "right": (2, 2), "drop": (2, 1),
    "pop": (53, 16), "zoom": (53, 16), "spin": (31, 0), "flip": (10, 0),
}

# a different transition per slide (PowerPoint 2010+ with a fade fallback)
TRANSITIONS = [
    '<p14:vortex dir="r"/>', '<p14:gallery dir="l"/>', '<p14:conveyor dir="l"/>',
    '<p14:switch dir="r"/>', '<p14:flip dir="l"/>', '<p14:prism isContent="1"/>',
]

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
}
NSDECL = " ".join(f'xmlns:{k}="{v}"' for k, v in NS.items())


class Ids:
    def __init__(self):
        self.n = 2  # 1 = tmRoot, 2 = mainSeq

    def __call__(self):
        self.n += 1
        return self.n


def tgt(spid):
    return f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'


def prop_anim(nid, spid, attr, points, dur, ease=True):
    tavs = "".join(f'<p:tav tm="{tm}"><p:val><p:strVal val="{v}"/></p:val></p:tav>' for tm, v in points)
    easing = ' decel="100000"' if ease else ""
    return (
        f'<p:anim calcmode="lin" valueType="num"><p:cBhvr additive="base">'
        f'<p:cTn id="{nid()}" dur="{dur}" fill="hold"{easing}/>{tgt(spid)}'
        f'<p:attrNameLst><p:attrName>{attr}</p:attrName></p:attrNameLst></p:cBhvr>'
        f"<p:tavLst>{tavs}</p:tavLst></p:anim>"
    )


def rot_anim(nid, spid, start, end, dur):
    return (
        f'<p:anim calcmode="lin" valueType="num"><p:cBhvr>'
        f'<p:cTn id="{nid()}" dur="{dur}" fill="hold" decel="100000"/>{tgt(spid)}'
        f"<p:attrNameLst><p:attrName>style.rotation</p:attrName></p:attrNameLst></p:cBhvr>"
        f'<p:tavLst><p:tav tm="0"><p:val><p:fltVal val="{start}"/></p:val></p:tav>'
        f'<p:tav tm="100000"><p:val><p:fltVal val="{end}"/></p:val></p:tav></p:tavLst></p:anim>'
    )


def effect(nid, spid, kind, delay_ms, rot=0.0):
    preset, sub = PRESETS.get(kind, (10, 0))
    dur = 800
    fade = (
        f'<p:animEffect transition="in" filter="fade"><p:cBhvr>'
        f'<p:cTn id="{nid()}" dur="{min(dur, 500)}"/>{tgt(spid)}</p:cBhvr></p:animEffect>'
    )
    b = [fade]
    if kind == "up":
        b.append(prop_anim(nid, spid, "ppt_y", [(0, "#ppt_y+0.06"), (100000, "#ppt_y")], dur))
    elif kind == "left":
        b.append(prop_anim(nid, spid, "ppt_x", [(0, "#ppt_x-0.06"), (100000, "#ppt_x")], dur))
    elif kind == "right":
        b.append(prop_anim(nid, spid, "ppt_x", [(0, "#ppt_x+0.06"), (100000, "#ppt_x")], dur))
    elif kind == "drop":
        b.append(prop_anim(nid, spid, "ppt_y",
                           [(0, "#ppt_y-0.12"), (60000, "#ppt_y+0.012"), (80000, "#ppt_y-0.005"), (100000, "#ppt_y")],
                           dur, ease=False))
    elif kind in ("pop", "zoom", "spin"):
        pts_w = [(0, "0"), (70000, "#ppt_w*1.07"), (100000, "#ppt_w")] if kind != "zoom" else [(0, "0"), (100000, "#ppt_w")]
        pts_h = [(0, "0"), (70000, "#ppt_h*1.07"), (100000, "#ppt_h")] if kind != "zoom" else [(0, "0"), (100000, "#ppt_h")]
        b.append(prop_anim(nid, spid, "ppt_w", pts_w, dur, ease=False))
        b.append(prop_anim(nid, spid, "ppt_h", pts_h, dur, ease=False))
        if kind == "spin":
            b.append(rot_anim(nid, spid, rot - 200, rot, dur))
    elif kind == "flip":
        b.append(prop_anim(nid, spid, "ppt_h", [(0, "0"), (100000, "#ppt_h")], dur))
        b.append(prop_anim(nid, spid, "ppt_y", [(0, "#ppt_y-0.03"), (100000, "#ppt_y")], dur))

    set_vis = (
        f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f"</p:cTn>{tgt(spid)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
        f'<p:to><p:strVal val="visible"/></p:to></p:set>'
    )
    return (
        f'<p:par><p:cTn id="{nid()}" presetID="{preset}" presetClass="entr" presetSubtype="{sub}" fill="hold" '
        f'nodeType="withEffect"><p:stCondLst><p:cond delay="{delay_ms}"/></p:stCondLst>'
        f"<p:childTnLst>{set_vis}{''.join(b)}</p:childTnLst></p:cTn></p:par>"
    )


def timing_xml(anims):
    nid = Ids()
    outer = nid()
    inner = nid()
    effects = "".join(effect(nid, *a) for a in anims)
    return (
        f"<p:timing {NSDECL}><p:tnLst><p:par>"
        f'<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
        f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="{inner}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f"<p:childTnLst>{effects}</p:childTnLst></p:cTn></p:par>"
        f"</p:childTnLst></p:cTn></p:par>"
        f"</p:childTnLst></p:cTn>"
        f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
        f"</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
    )


def transition_xml(i):
    return (
        f"<mc:AlternateContent {NSDECL}>"
        f'<mc:Choice Requires="p14"><p:transition spd="slow" p14:dur="1400">{TRANSITIONS[i % len(TRANSITIONS)]}</p:transition></mc:Choice>'
        f'<mc:Fallback><p:transition spd="slow"><p:fade/></p:transition></mc:Fallback>'
        f"</mc:AlternateContent>"
    )


def gif_size(path):
    with open(path, "rb") as f:
        head = f.read(10)
    return struct.unpack("<HH", head[6:10])


def add_gif(slide, it):
    x, y, w, h = (Emu(int(v * PX)) for v in (it["x"], it["y"], it["w"], it["h"]))
    ring_px = 6
    ring = slide.shapes.add_shape(MSO_SHAPE.OVAL, x - Emu(ring_px * PX), y - Emu(ring_px * PX),
                                  w + Emu(2 * ring_px * PX), h + Emu(2 * ring_px * PX))
    ring.fill.solid()
    ring.fill.fore_color.rgb = RGBColor(*it["ring"])
    ring.line.fill.background()
    ring.shadow.inherit = False

    pic = slide.shapes.add_picture(it["gif"], x, y, w, h)
    gw, gh = gif_size(it["gif"])
    if gw > gh:  # centre-crop to a square so the circle isn't squashed
        c = (1 - gh / gw) / 2
        pic.crop_left = pic.crop_right = c
    elif gh > gw:
        c = (1 - gw / gh) / 2
        pic.crop_top = pic.crop_bottom = c
    pic.auto_shape_type = MSO_SHAPE.OVAL
    pic.line.color.rgb = RGBColor(0xFF, 0xF8, 0xF0)
    pic.line.width = Pt(it["border"] * 0.75)

    # group ring + picture so they rotate/animate together
    grp = slide.shapes.add_group_shape([ring, pic])
    grp.rotation = it["rot"]
    return grp


def main():
    manifest = json.loads((BUILD / "manifest.json").read_text())
    prs = Presentation()
    prs.slide_width = Emu(1920 * PX)
    prs.slide_height = Emu(1080 * PX)
    blank = prs.slide_layouts[6]

    for i, items in enumerate(manifest["slides"]):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(manifest["bg"], 0, 0, prs.slide_width, prs.slide_height)

        anims = []
        for it in sorted(items, key=lambda t: t["delay"]):
            if "gif" in it:
                shape = add_gif(slide, it)
                rot = it["rot"]
            else:
                shape = slide.shapes.add_picture(it["png"], Emu(int(it["x"] * PX)), Emu(int(it["y"] * PX)),
                                                 Emu(int(it["w"] * PX)), Emu(int(it["h"] * PX)))
                rot = 0
            anims.append((shape.shape_id, it["anim"], int(it["delay"] * 1000), rot))

        sld = slide._element
        sld.append(etree.fromstring(transition_xml(i)))
        sld.append(etree.fromstring(timing_xml(anims)))

        slide.notes_slide.notes_text_frame.text = NOTES[i] if i < len(NOTES) else ""

    prs.save(OUT)
    print(f"saved {OUT} ({OUT.stat().st_size / 1e6:.1f} MB)")


if __name__ == "__main__":
    main()
